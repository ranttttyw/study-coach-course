"""仅解析文字型 PPTX；图片、扫描文字和复杂图表不在此 MVP 的提取范围内。"""
from pathlib import Path


def parse_pptx(path, max_slides=100):
    try:
        from pptx import Presentation
    except ImportError as exc:
        raise RuntimeError("缺少 python-pptx，请先重新同步项目依赖") from exc
    presentation = Presentation(path)
    if len(presentation.slides) > max_slides:
        raise ValueError(f"课件页数超过上限（{max_slides} 页）")
    slides = []
    for number, slide in enumerate(presentation.slides, 1):
        parts, table_parts = [], []
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False) and shape.text.strip():
                parts.append(shape.text.strip())
            if getattr(shape, "has_table", False):
                for row in shape.table.rows:
                    cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if cells:
                        table_parts.append(" | ".join(cells))
        notes = ""
        try:
            frame = slide.notes_slide.notes_text_frame
            notes = frame.text.strip() if frame else ""
        except Exception:
            pass
        title = ""
        try:
            title = slide.shapes.title.text.strip() if slide.shapes.title else ""
        except Exception:
            pass
        body = "\n".join(dict.fromkeys(parts + table_parts))
        slides.append({"slide_number": number, "title": title, "body_text": body, "speaker_notes": notes})
    return slides
