import tempfile
import unittest
from pathlib import Path

from auth import AuthStore
from knowledge_store import KnowledgeStore
from memory_utils import MemoryStore
from ppt_parser import parse_pptx
from pdf_parser import parse_pdf
from llm_gateway import CoachModel
from config import Settings
from outline_generator import create_chapter_outline


class MvpTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.path = Path(self.tmp.name) / "test.db"

    def tearDown(self):
        self.tmp.cleanup()

    def test_local_account_uses_server_generated_identity(self):
        auth = AuthStore(self.path)
        user = auth.register("student.one", "correct horse battery staple")
        self.assertNotEqual(user["id"], user["username"])
        self.assertEqual(auth.authenticate("student.one", "correct horse battery staple"), user)
        self.assertIsNone(auth.authenticate("student.one", "wrong password"))

    def test_attempt_survives_new_store_and_cannot_be_resubmitted(self):
        store = MemoryStore(self.path, "alice")
        attempt = {"id": "attempt-1", "status": "draft", "log": {"topic": "Memory"},
                   "profile": {}, "quiz": {"quiz": []}, "answers": [], "explanation": "", "result": None,
                   "mode": "demo", "model": "demo"}
        store.save_attempt(attempt)
        self.assertEqual(MemoryStore(self.path, "alice").latest_draft()["id"], "attempt-1")
        graded = store.mark_attempt_graded("attempt-1", {"score": None}, ["answer"], "summary")
        self.assertEqual(graded["status"], "graded")
        with self.assertRaises(ValueError):
            store.mark_attempt_graded("attempt-1", {"score": None}, ["answer"], "summary")

    def test_course_material_search_has_page_sources_and_user_scope(self):
        alice = KnowledgeStore(self.path, "alice")
        document_id = alice.create_document("memory.pptx", b"pptx bytes", "/tmp/memory.pptx", 10)
        alice.mark_ready(document_id, [{"slide_number": 8, "title": "长期记忆", "body_text": "语义记忆保存稳定知识", "speaker_notes": "课堂重点"}])
        result = alice.search("长期记忆是什么")
        self.assertEqual(result[0]["slide_number"], 8)
        self.assertEqual(result[0]["source"], "《memory.pptx》第 8 页")
        self.assertEqual(KnowledgeStore(self.path, "bob").search("长期记忆"), [])
        alice.clear_user()
        self.assertEqual(alice.search("长期记忆"), [])

    def test_pptx_parser_extracts_title_text_table_and_notes(self):
        from pptx import Presentation
        from pptx.util import Inches

        path = Path(self.tmp.name) / "lecture.pptx"
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[5])
        slide.shapes.title.text = "长期记忆"
        slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1)).text = "语义记忆保存稳定知识"
        table = slide.shapes.add_table(1, 2, Inches(1), Inches(2), Inches(5), Inches(1)).table
        table.cell(0, 0).text = "类型"
        table.cell(0, 1).text = "情景记忆"
        slide.notes_slide.notes_text_frame.text = "讲师备注"
        presentation.save(path)
        parsed = parse_pptx(path)
        self.assertEqual(parsed[0]["title"], "长期记忆")
        self.assertIn("语义记忆", parsed[0]["body_text"])
        self.assertIn("情景记忆", parsed[0]["body_text"])
        self.assertEqual(parsed[0]["speaker_notes"], "讲师备注")

    def test_pdf_parser_rejects_pdf_without_extractable_text(self):
        from pypdf import PdfWriter

        path = Path(self.tmp.name) / "scan.pdf"
        writer = PdfWriter()
        writer.add_blank_page(width=300, height=300)
        with path.open("wb") as stream:
            writer.write(stream)
        with self.assertRaisesRegex(ValueError, "OCR"):
            parse_pdf(path)

    def test_goals_chapters_and_materials_are_scoped(self):
        account = MemoryStore(self.path, "alice")
        first = account.create_goal("学习记忆", 4, "2026-10-24", 6, ["Memory"])
        second = account.create_goal("学习统计", 6, "2026-10-24", 4, ["Statistics"])
        first_store = MemoryStore(self.path, "alice", first["id"])
        first_store.add_chapter("工作记忆", 1)
        chapter = first_store.list_chapters()[0]
        first_store.complete_chapter(chapter["id"])
        self.assertEqual(first_store.list_chapters()[0]["status"], "done")
        self.assertEqual(MemoryStore(self.path, "alice", second["id"]).list_chapters(), [])
        materials = KnowledgeStore(self.path, "alice", first["id"])
        doc = materials.create_document("memory.pptx", b"memory", "/tmp/memory.pptx", 10)
        materials.mark_ready(doc, [{"slide_number": 1, "title": "工作记忆", "body_text": "临时上下文", "speaker_notes": ""}])
        self.assertTrue(materials.search("工作记忆"))
        self.assertEqual(KnowledgeStore(self.path, "alice", second["id"]).search("工作记忆"), [])

    def test_outline_generator_and_detected_chapters(self):
        account = MemoryStore(self.path, "alice")
        goal = account.create_goal("学习记忆", 4, "2026-10-24", 6, ["Memory"])
        store = MemoryStore(self.path, "alice", goal["id"])
        outline = create_chapter_outline([
            {"slide_number": 1, "title": "工作记忆", "body_text": "临时上下文", "speaker_notes": ""},
            {"slide_number": 2, "title": "长期记忆", "body_text": "跨会话事实", "speaker_notes": ""},
        ], 4, CoachModel(Settings(mode="demo")))
        store.add_detected_chapters("doc-1", outline)
        chapters = store.list_chapters()
        self.assertEqual([row["title"] for row in chapters], ["工作记忆", "长期记忆"])
        self.assertEqual(chapters[0]["start_slide"], 1)

    def test_new_material_merges_knowledge_without_resetting_completion(self):
        account = MemoryStore(self.path, "alice")
        goal = account.create_goal("学习记忆", 4, "2026-10-24", 6, ["Memory"])
        store = MemoryStore(self.path, "alice", goal["id"])
        store.merge_detected_chapters([{"title": "工作记忆", "document_id": "doc-1", "start_slide": 1, "end_slide": 3}])
        first = store.list_chapters()[0]
        store.complete_chapter(first["id"])
        result = store.merge_detected_chapters([
            {"title": "工作记忆", "document_id": "doc-2", "start_slide": 2, "end_slide": 5},
            {"title": "长期记忆", "document_id": "doc-2", "start_slide": 6, "end_slide": 9},
        ])
        chapters = store.list_chapters()
        self.assertEqual(result["added"], 1)
        self.assertEqual(len(chapters), 2)
        self.assertEqual(chapters[0]["status"], "done")

    def test_completed_goal_is_sorted_last_and_goal_deletion_is_scoped(self):
        account = MemoryStore(self.path, "alice")
        first = account.create_goal("已完成", 4, "2026-10-24", 6, ["Memory"])
        second = account.create_goal("进行中", 4, "2026-10-24", 6, ["AI"])
        first_store = MemoryStore(self.path, "alice", first["id"])
        first_store.add_chapter("知识点", 1)
        first_store.complete_chapter(first_store.list_chapters()[0]["id"])
        goals = account.list_goals()
        self.assertEqual([goal["title"] for goal in goals], ["进行中", "已完成"])
        self.assertTrue(account.delete_goal(first["id"]))
        self.assertEqual([goal["id"] for goal in account.list_goals()], [second["id"]])

    def test_knowledge_requires_confirmation_and_assessment_updates_mastery(self):
        account = MemoryStore(self.path, "alice")
        goal = account.create_goal("学习网络", 4, "2026-10-24", 6, ["Networking"])
        store = MemoryStore(self.path, "alice", goal["id"])
        store.merge_detected_chapters([{"title": "协议分层", "document_id": "doc-1", "start_slide": 1, "end_slide": 4}])
        self.assertEqual(store.list_chapters(confirmed_only=True), [])
        point = store.list_chapters()[0]
        store.update_knowledge_point(point["id"], "协议分层", "应用", 60, True)
        store.confirm_knowledge_plan()
        confirmed = store.list_chapters(confirmed_only=True)[0]
        self.assertEqual(confirmed["learning_level"], "应用")
        mastery = store.update_mastery(confirmed["id"], 95, [{"mastery": "掌握"}] * 3)
        self.assertEqual(mastery, "已掌握")
        self.assertEqual(store.list_chapters(confirmed_only=True)[0]["status"], "done")


if __name__ == "__main__":
    unittest.main()
